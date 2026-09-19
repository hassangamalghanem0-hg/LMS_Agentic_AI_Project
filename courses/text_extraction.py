"""Extracts plain text out of an uploaded Material file so the Instructor
Agent's AI tools (generate_summary_from_material / generate_quiz_from_material)
have something to read. Supports PDF, DOCX, PPTX and plain text; anything
else is marked 'unsupported' rather than failing loudly.
"""
import os

MAX_CHARS = 40_000  # keep prompts/summaries bounded


def guess_kind(filename):
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if ext == "pdf":
        return "pdf"
    if ext in ("docx",):
        return "docx"
    if ext in ("pptx",):
        return "pptx"
    if ext in ("txt", "md"):
        return "txt"
    if ext in ("mp4", "webm", "mov", "mkv", "avi", "m4v"):
        return "video"
    return "other"


def extract_text(file_field, kind):
    """file_field: a Django FieldFile (already saved to storage).
    Returns (text, status) where status is one of Material.ExtractionStatus values.
    Video files are intentionally never processed for text/AI -- they are
    plain uploaded lecture recordings, played back as-is.
    """
    if kind == "video":
        return "", "n/a"

    try:
        if kind == "txt":
            with file_field.open("rb") as f:
                raw = f.read()
            text = raw.decode("utf-8", errors="ignore")
            return text[:MAX_CHARS], "done"

        if kind == "pdf":
            from pypdf import PdfReader
            with file_field.open("rb") as f:
                reader = PdfReader(f)
                pages = [p.extract_text() or "" for p in reader.pages]
            text = "\n".join(pages).strip()
            return text[:MAX_CHARS], ("done" if text else "unsupported")

        if kind == "docx":
            import docx
            with file_field.open("rb") as f:
                doc = docx.Document(f)
                text = "\n".join(p.text for p in doc.paragraphs)
            return text[:MAX_CHARS], ("done" if text.strip() else "unsupported")

        if kind == "pptx":
            from pptx import Presentation
            with file_field.open("rb") as f:
                prs = Presentation(f)
                chunks = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text:
                            chunks.append(shape.text)
            text = "\n".join(chunks)
            return text[:MAX_CHARS], ("done" if text.strip() else "unsupported")

        return "", "unsupported"
    except Exception:
        return "", "failed"
